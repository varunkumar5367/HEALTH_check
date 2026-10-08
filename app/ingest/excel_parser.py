import openpyxl
from pathlib import Path
from typing import Dict, Any, List

class CMGKPIExcelParser:
    """
    Programmatic Excel parser for Nokia CMG KPI formulas and capacity thresholds workbook.
    """
    def __init__(self, excel_path: Path):
        self.excel_path = Path(excel_path)
        
    def parse_formulas_and_thresholds(self) -> Dict[str, Any]:
        result = {
            "node_sheets": {},
            "commercial_capacities": {}
        }
        
        if not self.excel_path.exists():
            return result
            
        try:
            wb = openpyxl.load_workbook(self.excel_path, data_only=True)
            for sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
                if sheet_name.startswith("KPI_FORMULAS"):
                    formulas = []
                    # Read rows from row 3 onwards
                    for r in range(3, sheet.max_row + 1):
                        kpi_name = sheet.cell(row=r, column=3).value
                        kpi_formula = sheet.cell(row=r, column=4).value
                        formula_type = sheet.cell(row=r, column=5).value
                        measurement = sheet.cell(row=r, column=6).value
                        if kpi_name:
                            formulas.append({
                                "kpi_name": kpi_name,
                                "formula": kpi_formula,
                                "formula_type": formula_type,
                                "measurement": measurement
                            })
                    result["node_sheets"][sheet_name] = formulas
                elif sheet_name == "COMMERCIAL":
                    for r in range(3, sheet.max_row + 1):
                        node_type = sheet.cell(row=r, column=2).value
                        kpi_name = sheet.cell(row=r, column=3).value
                        hw_cap = sheet.cell(row=r, column=6).value
                        sw_cap = sheet.cell(row=r, column=7).value
                        threshold = sheet.cell(row=r, column=8).value
                        if kpi_name:
                            result["commercial_capacities"][kpi_name] = {
                                "node_type": node_type,
                                "hw_capacity": hw_cap,
                                "sw_capacity": sw_cap,
                                "threshold": threshold
                            }
        except Exception as e:
            print(f"Error reading CMG Excel workbook {self.excel_path}: {e}")
            
        return result
